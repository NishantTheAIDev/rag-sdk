"""Loading source documents from disk.

Phase 2 ships a minimal loader supporting ``.txt``, ``.md``, and ``.json``
corpora. Phase 5 adds full document ingestion (PDF, DOCX, HTML) with
pluggable loader registry.
"""

from __future__ import annotations

import json
from abc import ABC, abstractmethod
from pathlib import Path
from typing import TYPE_CHECKING

from rag_sdk.config import DocumentLoaderConfig
from rag_sdk.core import Document, DocumentMetadata, Registry

if TYPE_CHECKING:
    pass


class IngestionError(ValueError):
    """Raised when a corpus cannot be loaded."""


class DocumentLoader(ABC):
    """Abstract base class for document loaders."""

    @abstractmethod
    def load(self, path: Path, config: DocumentLoaderConfig | None = None) -> list[Document]:
        """Load documents from a path (file or directory)."""


class TextLoader(DocumentLoader):
    """Load plain text and Markdown files."""

    def load(self, path: Path, config: DocumentLoaderConfig | None = None) -> list[Document]:
        if path.is_file():
            return [_load_text_file(path)]
        documents: list[Document] = []
        for file in sorted(path.iterdir()):
            if file.suffix.lower() not in {".txt", ".md"}:
                continue
            if not file.is_file():
                continue
            documents.append(_load_text_file(file))
        if not documents:
            raise IngestionError(f"No .txt or .md files found in {path}")
        return documents


class JSONLoader(DocumentLoader):
    """Load documents from a JSON file."""

    def load(self, path: Path, config: DocumentLoaderConfig | None = None) -> list[Document]:
        if not path.is_file() or path.suffix.lower() != ".json":
            raise IngestionError(f"JSON loader requires a .json file: {path}")
        return _load_json(path)


class PyPDFLoader(DocumentLoader):
    """Load documents from PDF files using pypdf."""

    def load(self, path: Path, config: DocumentLoaderConfig | None = None) -> list[Document]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise IngestionError(
                "pypdf is required for PDF loading. Install with: uv add 'rag-sdk[ingestion]'"
            ) from exc

        if not path.is_file() or path.suffix.lower() != ".pdf":
            raise IngestionError(f"PyPDF loader requires a .pdf file: {path}")

        try:
            reader = PdfReader(str(path))
        except Exception as exc:
            raise IngestionError(f"Could not read PDF: {path}") from exc

        documents: list[Document] = []
        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                continue
            doc_id = f"{path.stem}_page_{page_num}"
            documents.append(
                Document(
                    id=doc_id,
                    text=text,
                    metadata=DocumentMetadata(
                        source=str(path),
                        title=path.stem,
                        page=page_num,
                    ),
                )
            )
        if not documents:
            raise IngestionError(f"No extractable text found in PDF: {path}")
        return documents


class DocxLoader(DocumentLoader):
    """Load documents from DOCX files using python-docx."""

    def load(self, path: Path, config: DocumentLoaderConfig | None = None) -> list[Document]:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise IngestionError(
                "python-docx is required for DOCX loading. "
                "Install with: uv add 'rag-sdk[ingestion]'"
            ) from exc

        if not path.is_file() or path.suffix.lower() != ".docx":
            raise IngestionError(f"DOCX loader requires a .docx file: {path}")

        try:
            docx_doc = DocxDocument(str(path))
        except Exception as exc:
            raise IngestionError(f"Could not read DOCX: {path}") from exc

        text_parts: list[str] = []
        headings: list[str] = []
        for para in docx_doc.paragraphs:
            if para.style.name.startswith("Heading"):
                headings.append(para.text)
            text_parts.append(para.text)

        full_text = "\n".join(text_parts)
        if not full_text.strip():
            raise IngestionError(f"No extractable text found in DOCX: {path}")

        return [
            Document(
                id=path.stem,
                text=full_text,
                metadata=DocumentMetadata(
                    source=str(path),
                    title=path.stem,
                    headings=headings,
                ),
            )
        ]


class HTMLLoader(DocumentLoader):
    """Load documents from HTML files using BeautifulSoup + readability-lxml."""

    def load(self, path: Path, config: DocumentLoaderConfig | None = None) -> list[Document]:
        try:
            from bs4 import BeautifulSoup
        except ImportError as exc:
            raise IngestionError(
                "beautifulsoup4 is required for HTML loading. "
                "Install with: uv add 'rag-sdk[ingestion]'"
            ) from exc

        if not path.is_file() or path.suffix.lower() not in {".html", ".htm"}:
            raise IngestionError(f"HTML loader requires a .html or .htm file: {path}")

        try:
            html_content = path.read_text(encoding="utf-8")
        except Exception as exc:
            raise IngestionError(f"Could not read HTML file: {path}") from exc

        # Parse with BeautifulSoup first to extract metadata
        soup = BeautifulSoup(html_content, "lxml")

        # Extract title
        title = path.stem
        if soup.title and soup.title.string:
            title = soup.title.string.strip()

        # Extract headings from original HTML
        headings: list[str] = []
        for h in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
            text = h.get_text(strip=True)
            if text:
                headings.append(text)

        # Try readability-lxml for main content extraction
        main_text = ""
        try:
            from readability import Document as ReadabilityDocument
            read_doc = ReadabilityDocument(html_content)
            main_text = read_doc.summary()
            # Use readability title if available
            if read_doc.short_title():
                title = read_doc.short_title()
        except ImportError:
            pass  # fallback to BeautifulSoup

        if not main_text:
            # Fallback: use BeautifulSoup to extract text
            for elem in soup(["script", "style", "nav", "footer", "header", "aside"]):
                elem.decompose()
            main_text = soup.get_text(separator="\n", strip=True)

        if not main_text.strip():
            raise IngestionError(f"No extractable text found in HTML: {path}")

        return [
            Document(
                id=path.stem,
                text=main_text,
                metadata=DocumentMetadata(
                    source=str(path),
                    title=title,
                    headings=headings,
                ),
            )
        ]


# Registry for pluggable loaders
loader_registry: Registry[type[DocumentLoader]] = Registry()
loader_registry.register("text", TextLoader)
loader_registry.register("json", JSONLoader)
loader_registry.register("pypdf", PyPDFLoader)
loader_registry.register("docx", DocxLoader)
loader_registry.register("html", HTMLLoader)


def build_loader(config: DocumentLoaderConfig) -> DocumentLoader:
    """Resolve and construct a document loader from its configuration."""
    loader_type = loader_registry.get(config.strategy)
    return loader_type()


def load_documents(path: str | Path, config: DocumentLoaderConfig | None = None) -> list[Document]:
    """Load all documents from a directory or file using the configured loader.

    Backward-compatible default: TextLoader for .txt/.md, JSONLoader for .json.
    """
    source = Path(path)
    if not source.exists():
        raise IngestionError(f"Document path does not exist: {source}")

    if config is None:
        # Backward compatibility: auto-detect by extension
        if source.is_file():
            if source.suffix.lower() == ".json":
                return JSONLoader().load(source)
            if source.suffix.lower() == ".pdf":
                return PyPDFLoader().load(source)
            if source.suffix.lower() == ".docx":
                return DocxLoader().load(source)
            if source.suffix.lower() in {".html", ".htm"}:
                return HTMLLoader().load(source)
            return TextLoader().load(source)
        return TextLoader().load(source)

    loader = build_loader(config)
    return loader.load(source, config)


register_loader = loader_registry.decorator


def _load_text_file(path: Path) -> Document:
    try:
        text = path.read_text(encoding="utf-8")
    except OSError as exc:
        raise IngestionError(f"Could not read document: {path}") from exc
    return Document(
        id=path.stem,
        text=text,
        metadata=DocumentMetadata(source=str(path), title=path.stem),
    )


def _load_json(path: Path) -> list[Document]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except OSError as exc:
        raise IngestionError(f"Could not read document: {path}") from exc
    except json.JSONDecodeError as exc:
        raise IngestionError(f"Invalid JSON in {path}: {exc}") from exc
    if not isinstance(data, list):
        raise IngestionError(f"{path} must contain a list of documents")
    documents: list[Document] = []
    for index, entry in enumerate(data):
        if not isinstance(entry, dict):
            raise IngestionError(f"{path}: entry {index} must be an object")
        text = entry.get("text")
        if not isinstance(text, str) or not text:
            raise IngestionError(f"{path}: entry {index} must have non-empty 'text'")
        doc_id = entry.get("id")
        if doc_id is None:
            doc_id = f"doc-{index}"
        if not isinstance(doc_id, str):
            raise IngestionError(f"{path}: entry {index} 'id' must be a string")
        metadata = entry.get("metadata", {})
        if not isinstance(metadata, dict):
            raise IngestionError(f"{path}: entry {index} 'metadata' must be an object")
        documents.append(
            Document(
                id=doc_id,
                text=text,
                metadata=DocumentMetadata.model_validate(metadata),
            )
        )
    if not documents:
        raise IngestionError(f"No documents found in {path}")
    return documents